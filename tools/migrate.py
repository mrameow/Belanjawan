"""Tukar helaian lama (2025 + 2026) kepada senarai transaksi yang bersih.

Guna:  python3 -I tools/migrate.py <Belanjawan_Kelab_Guru.xlsx> <output.json>

Satu baris sumber = satu transaksi: Tarikh, Ahli, Butiran, Kategori,
Bulan (bulan yuran yang dibayar, cth "2026-03, 2026-04"), Kelab, Lawatan, Lain, Keluar.
Nombor dalam kurungan, cth "Amir (3, 4)", ialah BULAN yang dibayar.
  "12[30]"            -> bulan 12 dibayar separa RM30   (token "2025-12:30")
  "setengah bulan 12" -> separuh kadar bulan 12
"""
import datetime as dt
import json
import re
import sys

import openpyxl

AHLI = ["Amir", "Asnawi", "Bang Mie", "Boss Khairil", "Boss Wan", "Chang", "Hadi",
        "Hasnol", "Joe", "Khairiri", "Kuchai", "Suhaimi", "Ustaz", "Wafi", "Aiman"]
ALIAS = {"Boss": "Boss Wan"}  # helaian 2025 tulis "Boss (6)"
RATE = {"kelab": 35, "lawatan": 50}
START = (2025, 6)  # bulan pertama rekod


def to_date(v):
    """Excel menukar 'dd/mm' kepada mm/dd bila dd <= 12 -> tukar balik."""
    if isinstance(v, dt.datetime):
        return dt.date(v.year, v.day, v.month)
    d, m, y = v.strip().split("/")
    return dt.date(int(y), int(m), int(d))


def find_ahli(text):
    t = text.strip()
    m = re.match(r"^([A-Za-z ]+?)\s*(\(|$)", t)
    head = ALIAS.get(m.group(1).strip(), m.group(1).strip()) if m else t
    if head in AHLI:
        return head
    m = re.search(r"-\s*([A-Za-z ]+?)\s*(\(|$)", t)  # "Majlis KKGS - Hasnol (Lamb Chop)"
    if m:
        tail = ALIAS.get(m.group(1).strip(), m.group(1).strip())
        if tail in AHLI:
            return tail
    return ""


def kategori_masuk(text):
    low = text.lower()
    if low.startswith("majlis"):
        return "Majlis"
    if "sumbangan" in low:
        return "Sumbangan"
    return "Lain-lain"


def kategori_keluar(text):
    low = text.lower()
    if low.startswith("refund"):
        return "Refund"
    if "pantri" in low or "filter air" in low:
        return "Duit Pantri"
    if "lawatan" in low or "trip" in low or "deposit" in low:
        return "Lawatan / Trip"
    if "baju" in low:
        return "Pakaian"
    if "pskpp" in low or "buku perancangan" in low:
        return "Langganan / Bahan"
    if any(k in low for k in ("majlis", "hadiah", "jamuan", "hari guru", "hamper", "boling",
                              "kahwin", "sumbangan", "minuman", "makanan")):
        return "Majlis & Hadiah"
    return "Pelbagai"


def month_tokens(butiran, rate):
    """[(bulan, amaun_separa|None)] daripada kurungan bulan; [] jika tiada."""
    for grp in re.findall(r"\(([^)]*)\)", butiran):
        rest = re.sub(r"kelab|lawatan|yuran|setengah\s+bulan|[\d\s,&\[\]]", "", grp, flags=re.I)
        if rest or not re.search(r"\d", grp):
            continue  # bukan kurungan bulan (cth "Lamb Chop", "Hutang Lawatan")
        out, g = [], grp
        half = re.search(r"setengah\s+bulan\s+(\d+)", g, re.I)
        if half:
            out.append((int(half.group(1)), rate / 2))
            g = g[:half.start()] + g[half.end():]
        g = re.sub(r"kelab|lawatan|yuran|&", " ", g, flags=re.I)
        out = [(int(n), float(a) if a else None) for n, a in re.findall(r"(\d+)(?:\[(\d+)\])?", g)] + out
        return out
    return []


def next_free(claimed, start=START):
    y, m = start
    while (y, m) in claimed:
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return y, m


def assign_year(m, claimed_sets):
    """Bulan >= 6 -> cuba 2025 dulu; jika sudah dituntut oleh ahli itu -> 2026, dst."""
    y = 2025 if m >= 6 else 2026
    while any((y, m) in c for c in claimed_sets):
        y += 1
    return y


def parse(ws):
    rows, cur, in_yuran = [], None, False
    for r in range(2, ws.max_row + 1):
        a, b, c, d, e = (ws.cell(r, i).value for i in range(1, 6))
        if a is not None:
            cur, in_yuran = to_date(a), False
        if b is None:
            continue
        if b == "Yuran" and c == "Kelab":
            in_yuran = True
            continue
        if b == "Baki daripada 2025":
            continue  # 2025 dimasukkan penuh, jadi baki bawa-ke-hadapan tidak perlu
        b = str(b).strip()
        kelab = float(c) if isinstance(c, (int, float)) else 0.0
        lawatan = float(d) if isinstance(d, (int, float)) else 0.0
        keluar = float(e) if isinstance(e, (int, float)) else 0.0
        ahli = find_ahli(b)
        base = dict(tarikh=cur, ahli=ahli, butiran=b, bulan="", kelab=0.0, lawatan=0.0, lain=0.0, keluar=0.0)
        if keluar:
            rows.append({**base, "kategori": kategori_keluar(b), "keluar": keluar})
        if in_yuran and ahli and not b.lower().startswith(("majlis", "sumbangan")):
            if kelab or lawatan:
                rows.append({**base, "kategori": "Yuran", "kelab": kelab, "lawatan": lawatan, "_raw": True})
        elif kelab or lawatan:
            rows.append({**base, "kategori": kategori_masuk(b), "lain": kelab + lawatan})
    return rows


def fill_months(rows):
    """Isi lajur Bulan. Mengikut tarikh bayaran supaya tuntutan bulan berurutan."""
    claimed = {(a, k): set() for a in AHLI for k in ("kelab", "lawatan")}
    inferred = []
    for x in sorted((x for x in rows if x.get("_raw")), key=lambda x: x["tarikh"]):
        cats = [k for k in ("kelab", "lawatan") if x[k]]
        rate = RATE[cats[0]]
        toks = month_tokens(x["butiran"], rate)
        out = []
        if toks:
            for m, amt in toks:
                y = assign_year(m, [claimed[(x["ahli"], k)] for k in cats])
                out.append((y, m, amt))
        else:  # tiada bulan ditulis -> isi bulan kosong terawal ikut amaun
            amount = x[cats[0]]
            n, rem = int(amount // rate), amount % rate
            cl = claimed[(x["ahli"], cats[0])]
            for _ in range(n):
                y, m = next_free(cl | {(o[0], o[1]) for o in out})
                out.append((y, m, None))
            if rem:
                y, m = next_free(cl | {(o[0], o[1]) for o in out})
                out.append((y, m, rem))
            inferred.append((x["tarikh"].isoformat(), x["butiran"], cats, [f"{o[0]}-{o[1]:02d}" for o in out]))
        for k in cats:
            claimed[(x["ahli"], k)].update((y, m) for y, m, _ in out)
        x["bulan"] = ", ".join(f"{y}-{m:02d}" + (f":{a:g}" if a else "") for y, m, a in out)
    return inferred


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    rows = parse(wb["2025"]) + parse(wb["2026"])
    rows.sort(key=lambda x: x["tarikh"])  # sort stabil: susunan asal dalam hari yang sama kekal
    inferred = fill_months(rows)
    for x in rows:
        x.pop("_raw", None)
        x["tarikh"] = x["tarikh"].isoformat()
    masuk = sum(x["kelab"] + x["lawatan"] + x["lain"] for x in rows)
    keluar = sum(x["keluar"] for x in rows)
    print(f"{len(rows)} baris | masuk {masuk:.2f} | keluar {keluar:.2f} | baki {masuk - keluar:.2f}")
    print(f"{len(inferred)} baris yuran tiada bulan ditulis - bulan diisi automatik (semak):")
    for i in inferred:
        print("  ", *i)
    json.dump(rows, open(dst, "w"), ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main(*sys.argv[1:3])
