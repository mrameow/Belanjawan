"""Bina Belanjawan_Kelab_Guru_Baru.xlsx (import ke Google Sheets).

Guna:  python3 -I tools/build_sheet.py <tx.json> <output.xlsx>
"""
import datetime as dt
import json
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

TEAL, TEAL_D, TEAL_L = "0F766E", "134E4A", "F0FDFA"
GREEN, GREEN_L, RED, RED_L = "15803D", "DCFCE7", "B91C1C", "FEE2E2"
GREY, GREY_L, INK = "64748B", "F1F5F9", "0F172A"
FONT = "Arial"
RM = '"RM"#,##0.00;[Red]-"RM"#,##0.00'
MAX_ROW = 3000

AHLI = ["Amir", "Asnawi", "Bang Mie", "Boss Khairil", "Boss Wan", "Chang", "Hadi",
        "Hasnol", "Joe", "Khairiri", "Kuchai", "Suhaimi", "Ustaz", "Wafi", "Aiman"]
# Mula / Tamat (YYYY, M) -- anggaran daripada data; boleh diubah di helaian Tetapan
TEMPOH = {
    "Boss Khairil": ((2025, 6), (2025, 9)), "Suhaimi": ((2025, 6), (2025, 8)),
    "Wafi": ((2025, 6), (2026, 3)), "Aiman": ((2025, 12), None),
}
NO_LAWATAN = {"Kuchai": "Tiada kutipan lawatan (refund 400 Jan 2026) - semak",
              "Wafi": "Lawatan diuruskan berasingan (Bayaran Lebih Lawatan) - semak"}
CAT_MASUK = ["Yuran Kelab", "Yuran Lawatan", "Majlis", "Sumbangan", "Lain-lain"]
CAT_KELUAR = ["Duit Pantri", "Lawatan / Trip", "Majlis & Hadiah", "Pakaian",
              "Langganan / Bahan", "Refund", "Pelbagai"]


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    kw.setdefault("color", INK)
    return Font(**kw)


thin = Side(style="thin", color="E2E8F0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOT = Border(bottom=thin)


def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = fill(TEAL_D)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26


def banner(ws, text, sub, last_col):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    a, b = ws.cell(1, 1, text), ws.cell(2, 1, sub)
    a.font, b.font = font(size=20, bold=True, color="FFFFFF"), font(size=10, color="CCFBF1")
    for r in (1, 2):
        for c in range(1, last_col + 1):
            ws.cell(r, c).fill = fill(TEAL)
        ws.cell(r, 1).alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 20


def build(tx, out):
    wb = Workbook()
    ws_r, ws_t, ws_y, ws_s = wb.active, wb.create_sheet("Transaksi"), wb.create_sheet("Yuran Ahli"), wb.create_sheet("Tetapan")
    ws_r.title = "Ringkasan"

    # ---------------- Tetapan ----------------
    banner(ws_s, "Tetapan", "Senarai ahli, kadar yuran dan kategori. Aplikasi membaca senarai ini.", 8)
    header(ws_s, 4, ["Ahli", "Mula", "Tamat", "Kelab / bulan", "Lawatan / bulan", "Aktif", "Catatan"])
    for i, n in enumerate(AHLI):
        r = 5 + i
        mula, tamat = TEMPOH.get(n, ((2025, 6), None))
        ws_s.cell(r, 1, n)
        ws_s.cell(r, 2, dt.date(*mula, 1))
        ws_s.cell(r, 3, dt.date(*tamat, 1) if tamat else None)
        ws_s.cell(r, 4, 35)
        ws_s.cell(r, 5, 0 if n in NO_LAWATAN else 50)
        ws_s.cell(r, 7, NO_LAWATAN.get(n))
        ws_s.cell(r, 6, "Tidak" if tamat else "Ya")
        for c in range(1, 8):
            cell = ws_s.cell(r, c)
            cell.font, cell.border = font(), BOX
            if c in (2, 3):
                cell.number_format = "mmm yyyy"
            if c in (4, 5):
                cell.number_format = RM
    ws_s.cell(21, 1, "Lawatan dikutip sehingga").font = font(bold=True)
    ws_s.cell(21, 2, dt.date(2026, 2, 1)).number_format = "mmm yyyy"
    ws_s.cell(22, 1, "Tarikh mula rekod").font = font(bold=True)
    ws_s.cell(22, 2, dt.date(2025, 6, 1)).number_format = "mmm yyyy"
    for r in (21, 22):
        ws_s.cell(r, 2).fill = fill("FEF9C3")
        ws_s.cell(r, 2).font = font()
    ws_s.cell(24, 1, "Kolum kuning / jadual ahli boleh diubah. Tarikh 'Mula' dan 'Tamat' hanyalah anggaran dari data lama - semak semula.").font = font(italic=True, color=GREY)
    header(ws_s, 4, ["Kategori Masuk"], col=9)
    header(ws_s, 4, ["Kategori Keluar"], col=10)
    for i, v in enumerate(CAT_MASUK):
        ws_s.cell(5 + i, 9, v).font = font()
    for i, v in enumerate(CAT_KELUAR):
        ws_s.cell(5 + i, 10, v).font = font()
    for col, w in zip("ABCDEFGHIJ", (22, 12, 12, 15, 16, 8, 52, 3, 20, 22)):
        ws_s.column_dimensions[col].width = w
    ws_s.freeze_panes = "A5"

    # ---------------- Transaksi ----------------
    ws_t.sheet_view.showGridLines = False
    header(ws_t, 1, ["ID", "Tarikh", "Jenis", "Kategori", "Ahli / Pihak", "Butiran", "Jumlah (RM)", "Baki (RM)"])
    for i, x in enumerate(tx):
        r = i + 2
        ws_t.cell(r, 1, f"M{i + 1:04d}")
        ws_t.cell(r, 2, dt.date.fromisoformat(x["tarikh"]))
        ws_t.cell(r, 3, x["jenis"])
        ws_t.cell(r, 4, x["kategori"])
        ws_t.cell(r, 5, x["ahli"])
        ws_t.cell(r, 6, x["butiran"])
        ws_t.cell(r, 7, x["jumlah"])
    n_last = len(tx) + 1
    for r in range(2, MAX_ROW + 1):
        ws_t.cell(r, 8, f'=IF(G{r}="","",SUMIF(C$2:C{r},"Masuk",G$2:G{r})-SUMIF(C$2:C{r},"Keluar",G$2:G{r}))')
        for c in range(1, 9):
            cell = ws_t.cell(r, c)
            cell.font = font(color=GREY if c in (1, 8) else INK)
            cell.border = BOT
            if r % 2 == 1:
                cell.fill = fill("F8FAFC")
        ws_t.cell(r, 2).number_format = "dd mmm yyyy"
        ws_t.cell(r, 2).alignment = Alignment(horizontal="left")
        ws_t.cell(r, 3).alignment = Alignment(horizontal="center")
        ws_t.cell(r, 7).number_format = RM
        ws_t.cell(r, 8).number_format = RM
    rng = f"C2:C{MAX_ROW}"
    ws_t.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Masuk"'], font=Font(name=FONT, bold=True, color=GREEN), fill=fill(GREEN_L)))
    ws_t.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Keluar"'], font=Font(name=FONT, bold=True, color=RED), fill=fill(RED_L)))
    ws_t.conditional_formatting.add(f"G2:G{MAX_ROW}", FormulaRule(formula=['$C2="Keluar"'], font=Font(name=FONT, color=RED)))
    ws_t.conditional_formatting.add(f"G2:G{MAX_ROW}", FormulaRule(formula=['$C2="Masuk"'], font=Font(name=FONT, color=GREEN)))
    dv_j = DataValidation(type="list", formula1='"Masuk,Keluar"', allow_blank=True)
    dv_k = DataValidation(type="list", formula1="=Tetapan!$I$5:$J$11", allow_blank=True, showErrorMessage=False)
    dv_a = DataValidation(type="list", formula1="=Tetapan!$A$5:$A$19", allow_blank=True, showErrorMessage=False)
    for dv, col in ((dv_j, "C"), (dv_k, "D"), (dv_a, "E")):
        ws_t.add_data_validation(dv)
        dv.add(f"{col}2:{col}{MAX_ROW}")
    for col, w in zip("ABCDEFGH", (8, 13, 9, 19, 16, 52, 14, 14)):
        ws_t.column_dimensions[col].width = w
    ws_t.freeze_panes = "A2"
    ws_t.auto_filter.ref = f"A1:H{MAX_ROW}"

    # ---------------- Ringkasan ----------------
    banner(ws_r, "Belanjawan Kelab Guru", "Ringkasan kewangan - dikemas kini automatik daripada helaian Transaksi", 12)
    T = "Transaksi!"
    kpis = [
        ("BAKI SEMASA", f'=SUMIF({T}C:C,"Masuk",{T}G:G)-SUMIF({T}C:C,"Keluar",{T}G:G)', TEAL),
        ("JUMLAH MASUK", f'=SUMIF({T}C:C,"Masuk",{T}G:G)', GREEN),
        ("JUMLAH KELUAR", f'=SUMIF({T}C:C,"Keluar",{T}G:G)', RED),
        ("BULAN INI (BERSIH)", f'=SUMIFS({T}G:G,{T}C:C,"Masuk",{T}B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))-SUMIFS({T}G:G,{T}C:C,"Keluar",{T}B:B,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))', "B45309"),
    ]
    for i, (label, f, colr) in enumerate(kpis):
        c0 = 1 + i * 3
        ws_r.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 1)
        ws_r.merge_cells(start_row=5, start_column=c0, end_row=5, end_column=c0 + 1)
        a, b = ws_r.cell(4, c0, label), ws_r.cell(5, c0, f)
        a.font, b.font = font(size=9, bold=True, color=GREY), font(size=18, bold=True, color=colr)
        b.number_format = RM
        for r in (4, 5):
            for c in (c0, c0 + 1):
                ws_r.cell(r, c).fill = fill(GREY_L)
        ws_r.cell(4, c0).alignment = Alignment(indent=1, vertical="bottom")
        ws_r.cell(5, c0).alignment = Alignment(indent=1, vertical="center", horizontal="left")
    ws_r.row_dimensions[4].height = 22
    ws_r.row_dimensions[5].height = 34

    ws_r.cell(7, 1, "Pergerakan bulanan").font = font(size=13, bold=True, color=TEAL_D)
    header(ws_r, 8, ["Bulan", "Masuk", "Keluar", "Bersih", "Baki akhir"])
    n_months = 19  # Jun 2025 - Dis 2026
    for i in range(n_months):
        r = 9 + i
        ws_r.cell(r, 1, f"=EDATE(Tetapan!$B$22,{i})").number_format = "mmm yyyy"
        crit = f'{T}B:B,">="&A{r},{T}B:B,"<"&EDATE(A{r},1)'
        ws_r.cell(r, 2, f'=SUMIFS({T}G:G,{T}C:C,"Masuk",{crit})')
        ws_r.cell(r, 3, f'=SUMIFS({T}G:G,{T}C:C,"Keluar",{crit})')
        ws_r.cell(r, 4, f"=B{r}-C{r}")
        ws_r.cell(r, 5, f"=SUM(D$9:D{r})")
        for c in range(1, 6):
            cell = ws_r.cell(r, c)
            cell.font, cell.border = font(), BOT
            if c > 1:
                cell.number_format = RM
        ws_r.cell(r, 2).font = font(color=GREEN)
        ws_r.cell(r, 3).font = font(color=RED)
        ws_r.cell(r, 5).font = font(bold=True)
    end = 8 + n_months
    ws_r.conditional_formatting.add(f"D9:D{end}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color=RED, bold=True)))

    ws_r.cell(7, 7, "Perbelanjaan mengikut kategori").font = font(size=13, bold=True, color=TEAL_D)
    header(ws_r, 8, ["Kategori", "Jumlah", "%"], col=7)
    for i, k in enumerate(CAT_KELUAR):
        r = 9 + i
        ws_r.cell(r, 7, f"=Tetapan!J{5 + i}")
        ws_r.cell(r, 8, f'=SUMIFS({T}G:G,{T}C:C,"Keluar",{T}D:D,G{r})').number_format = RM
        ws_r.cell(r, 9, f'=IF($C${end + 1}=0,0,H{r}/$C${end + 1})').number_format = "0%"
        for c in (7, 8, 9):
            ws_r.cell(r, c).font, ws_r.cell(r, c).border = font(), BOT
    ws_r.cell(end + 1, 1, "Jumlah").font = font(bold=True)
    for col in "BCD":
        c = ws_r[f"{col}{end + 1}"]
        c.value, c.number_format, c.font = f"=SUM({col}9:{col}{end})", RM, font(bold=True)
    for c in range(1, 6):
        ws_r.cell(end + 1, c).fill = fill(TEAL_L)

    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title, ch.height, ch.width = "Masuk vs Keluar", 8.5, 17
    ch.add_data(Reference(ws_r, min_col=2, max_col=3, min_row=8, max_row=end), titles_from_data=True)
    ch.set_categories(Reference(ws_r, min_col=1, min_row=9, max_row=end))
    ch.series[0].graphicalProperties.solidFill = GREEN
    ch.series[1].graphicalProperties.solidFill = RED
    ws_r.add_chart(ch, f"G{18}")
    ln = LineChart()
    ln.title, ln.height, ln.width = "Baki akhir", 7.5, 17
    ln.add_data(Reference(ws_r, min_col=5, min_row=8, max_row=end), titles_from_data=True)
    ln.set_categories(Reference(ws_r, min_col=1, min_row=9, max_row=end))
    ln.series[0].graphicalProperties.line.solidFill = TEAL
    ln.series[0].graphicalProperties.line.width = 28000
    ws_r.add_chart(ln, f"G{35}")
    for col, w in zip("ABCDEFGHI", (14, 15, 15, 15, 16, 3, 24, 15, 8)):
        ws_r.column_dimensions[col].width = w

    # ---------------- Yuran Ahli ----------------
    banner(ws_y, "Yuran Ahli", "Bayaran sebenar berbanding jangkaan (berdasarkan tempoh di Tetapan). Anggaran sahaja.", 10)
    header(ws_y, 4, ["Ahli", "Kelab dibayar", "Kelab jangkaan", "Tunggakan Kelab", "Lawatan dibayar", "Lawatan jangkaan", "Tunggakan Lawatan", "Jumlah tunggakan", "Status"])
    for i, n in enumerate(AHLI):
        r, s = 5 + i, 5 + i
        ws_y.cell(r, 1, f"=Tetapan!A{s}")
        ws_y.cell(r, 2, f'=SUMIFS({T}G:G,{T}C:C,"Masuk",{T}D:D,"Yuran Kelab",{T}E:E,A{r})')
        ws_y.cell(r, 3, f"=MAX(0,DATEDIF(Tetapan!B{s},IF(Tetapan!C{s}=\"\",TODAY(),MIN(Tetapan!C{s},TODAY())),\"M\")+1)*Tetapan!D{s}")
        ws_y.cell(r, 4, f"=MAX(0,C{r}-B{r})")
        ws_y.cell(r, 5, f'=SUMIFS({T}G:G,{T}C:C,"Masuk",{T}D:D,"Yuran Lawatan",{T}E:E,A{r})-SUMIFS({T}G:G,{T}C:C,"Keluar",{T}D:D,"Refund",{T}E:E,A{r})')
        ws_y.cell(r, 6, f"=MAX(0,DATEDIF(Tetapan!B{s},IF(Tetapan!C{s}=\"\",Tetapan!$B$21,MIN(Tetapan!C{s},Tetapan!$B$21)),\"M\")+1)*Tetapan!E{s}")
        ws_y.cell(r, 7, f"=MAX(0,F{r}-E{r})")
        ws_y.cell(r, 8, f"=D{r}+G{r}")
        ws_y.cell(r, 9, f'=IF(H{r}=0,"Selesai","Tunggak")')
        for c in range(1, 10):
            cell = ws_y.cell(r, c)
            cell.font, cell.border = font(bold=(c == 1)), BOX
            if 2 <= c <= 8:
                cell.number_format = RM
            if c == 9:
                cell.alignment = Alignment(horizontal="center")
    last = 4 + len(AHLI)
    ws_y.cell(last + 1, 1, "Jumlah").font = font(bold=True)
    for col in "BCDEFGH":
        c = ws_y[f"{col}{last + 1}"]
        c.value, c.number_format, c.font = f"=SUM({col}5:{col}{last})", RM, font(bold=True)
    for c in range(1, 10):
        ws_y.cell(last + 1, c).fill = fill(TEAL_L)
    ws_y.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator="equal", formula=['"Selesai"'], font=Font(name=FONT, bold=True, color=GREEN), fill=fill(GREEN_L)))
    ws_y.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator="equal", formula=['"Tunggak"'], font=Font(name=FONT, bold=True, color=RED), fill=fill(RED_L)))
    ws_y.conditional_formatting.add(f"H5:H{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, bold=True, color=RED)))
    for col, w in zip("ABCDEFGHI", (18, 15, 15, 16, 15, 16, 17, 16, 12)):
        ws_y.column_dimensions[col].width = w
    ws_y.freeze_panes = "B5"

    for ws, colr in ((ws_r, TEAL), (ws_t, "0EA5E9"), (ws_y, "F59E0B"), (ws_s, GREY)):
        ws.sheet_properties.tabColor = colr
    wb.save(out)


if __name__ == "__main__":
    build(json.load(open(sys.argv[1])), sys.argv[2])
